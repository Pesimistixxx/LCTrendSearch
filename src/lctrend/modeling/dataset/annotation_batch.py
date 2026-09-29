"""Expert annotation batch: historical slices plus the organisers' list.

The batch gives
people what the LLM labels cannot: a checked, independent set of labels.
It is a workbook of three sheets:

- «Правила»: how to label, in Russian;
- «Исторические срезы»: 300 technology × date rows, one per family, with
  quotas by period (older periods on the 36-month horizon, 2024–2025 on
  the 12-month one); within a quota about 60% of the rows have documents
  after T, the rest are a baseline. The LLM's view of the same year
  (verdict, score, hype, maturity, rationale) is shown as a hint to
  check, never as the answer;
- «Список 100»: the organisers' 100 weak signals of September 2026, with
  the graph technology each one was matched to (``labeling.benchmark``).

Selection is reproducible: families are ordered by a hash of the batch
name, so the same export gives the same batch. ``read_annotations``
reads a filled workbook back as expert labels.

    python -m lctrend.modeling.dataset.annotation_batch \\
        --review artifacts/modeling/R/dataset/review.reviewer_1.csv \\
        --llm artifacts/modeling/R/dataset/llm_labels.csv.jsonl \\
        --list outputs/annotation-batch-2026-09-29/batch.xlsx \\
        --output outputs/annotation-batch-R/batch.xlsx
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import defaultdict
from pathlib import Path
from typing import (
    Any,
    Callable,
    Dict,
    List,
    Mapping,
    Sequence,
    Tuple,
)

BATCH = "annotation-batch"
Quota = Tuple[str, int, Callable[[int, Mapping[str, str]], bool]]


def _complete(row, months):
    key = "calendar_complete_36m" if months == 36 else "calendar_complete_12m"
    return str(row.get(key)) == "True"


QUOTAS: Tuple[Quota, ...] = (
    ("1959–2010, 36m", 10, lambda y, r: y <= 2010 and _complete(r, 36)),
    (
        "2011–2016, 36m",
        20,
        lambda y, r: 2011 <= y <= 2016 and _complete(r, 36),
    ),
    (
        "2017–2020, 36m",
        100,
        lambda y, r: 2017 <= y <= 2020 and _complete(r, 36),
    ),
    (
        "2021–2023, 36m",
        110,
        lambda y, r: 2021 <= y <= 2023 and _complete(r, 36),
    ),
    ("2024, 12m", 30, lambda y, r: y == 2024 and _complete(r, 12)),
    ("2025, 12m", 30, lambda y, r: y == 2025 and _complete(r, 12)),
)
FOLLOWED_SHARE = 0.6

REVIEW_FIELDS = (
    "technology_id",
    "technology",
    "snapshot_date",
    "reviewer_state",
    "reviewer_signal_12m",
    "reviewer_signal_36m",
    "reviewer_trend_12m",
    "reviewer_trend_36m",
    "evidence_notes",
    "llm_verdict",
    "llm_score_at_t",
    "llm_hype_at_t",
    "llm_maturity_at_t",
    "llm_rationale",
    "family_id",
    "first_seen_date",
    "horizon_12m_end",
    "horizon_end",
    "calendar_complete_12m",
    "calendar_complete_36m",
    "document_count",
    "source_families",
    "organizations",
    "pre_t_sources",
    "followup_0_12m_sources",
    "followup_13_36m_sources",
    "sample_group",
)
REVIEWER_FIELDS = REVIEW_FIELDS[3:9]
LIST_FIELDS = (
    "source_no",
    "technology",
    "domain",
    "companies",
    "why_weak",
    "stage",
    "trend",
    "source_score",
    "source_urls",
    "graph_technology_id",
    "graph_technology",
    "match_cosine",
    "match_status",
    "match_notes",
)
STATES = ("weak", "trend", "mature", "faded", "insufficient", "rejected")

GUIDE = (
    (
        "Разметка слабых сигналов",
        "300 исторических срезов и 100 записей внешнего списка",
    ),
    (
        "Единица разметки",
        (
            "Одна технология на конкретную дату snapshot_date. Одна и та же "
            "технология может менять статус со временем."
        ),
    ),
    (
        "reviewer_state",
        (
            "weak — ранний достоверный сигнал на дату T; trend — уже заметный"
            " тренд; mature — зрелая технология; faded — не развилась; "
            "insufficient — не хватает данных; rejected — неверно выделенная "
            "технология (не технология)."
        ),
    ),
    (
        "reviewer_signal_12m / 36m",
        (
            "1 — на T это был слабый сигнал, впоследствии подтверждённый "
            "независимыми источниками в указанный срок. 0 — проверенный "
            "отрицательный исход. Пусто — неизвестно или охват источников "
            "недостаточен."
        ),
    ),
    (
        "reviewer_trend_12m / 36m",
        (
            "1 — за указанный срок технология стала заметным трендом; 0 — "
            "проверенный отрицательный исход; пусто — неизвестно. Не "
            "подменяйте слабый сигнал одним только ростом публикаций."
        ),
    ),
    (
        "Подсказка LLM",
        (
            "Колонки llm_* — мнение LLM о том же годе: вердикт траектории, "
            "оценка (0 — слабый сигнал, 1 — точно не слабый), перегретость, "
            "сформированность и обоснование. Это подсказка для проверки, а не"
            " ответ: если вы не согласны, ставьте свою метку и коротко "
            "объясните почему."
        ),
    ),
    (
        "Временная граница",
        (
            "Для оценки статуса на T смотрите только документы, датированные "
            "не позже T. Последующие документы нужны отдельно для проверки "
            "исхода за 12 или 36 месяцев."
        ),
    ),
    (
        "Неполный горизонт",
        (
            "Если calendar_complete_36m = FALSE, оставьте обе 36-месячные "
            "метки пустыми. Календарное завершение само по себе не доказывает"
            " полноту источников."
        ),
    ),
    (
        "evidence_notes",
        (
            "Укажите URL/ID решающих источников и коротко объясните 1 или 0. "
            "Если вывод невозможен, укажите, каких данных не хватает."
        ),
    ),
    (
        "Отбор строк",
        (
            "В партии 300 разных семейств технологий (после слияния дублей). "
            "Около 60% строк в каждой квоте отобраны по наличию последующих "
            "документов; партия не отражает долю слабых сигналов во всём "
            "графе."
        ),
    ),
    (
        "Ссылки в графе",
        (
            "В колонках источников — до четырёх документов за период. Это "
            "подсказки для проверки, а не полный обзор литературы. Пустая "
            "колонка не доказывает отсутствие результата."
        ),
    ),
    (
        "Ретроспектива",
        (
            "Срезы строились по датам документов из графа; документ мог быть "
            "получен системой позднее даты T. Для строгого backtest нужен "
            "аудит времени получения."
        ),
    ),
    (
        "Список 100",
        (
            "Кандидаты организаторов на сентябрь 2026, а не метки исхода. "
            "graph_technology — ближайшая по смыслу технология графа "
            "(эмбеддинги), match_cosine — сходство. Подтвердите match_status:"
            " same — та же технология; related — родственная; none — в графе "
            "нет."
        ),
    ),
)


def _rank(group: str, family: str) -> str:
    return hashlib.sha256(f"{BATCH}:{group}:{family}".encode()).hexdigest()


def _refs(row) -> List[Dict[str, Any]]:
    return json.loads(row.get("recent_documents") or "[]")


def show_sources(documents: Mapping[str, Mapping[str, Any]], limit=4) -> str:
    ordered = sorted(
        documents.values(),
        key=lambda item: (item["date"], item["document_id"]),
        reverse=True,
    )[:limit]
    lines = []
    for item in ordered:
        title = str(item.get("title") or item["document_id"])[:95]
        link = item.get("url") or item["document_id"]
        lines.append(f"{item['date']} | {title} | {link}")
    return "\n".join(lines)


def select_batch(
    rows: Sequence[Dict[str, str]],
    quotas: Sequence[Quota] = QUOTAS,
    exclude: Sequence[str] = (),
) -> List[Dict[str, Any]]:
    """One snapshot per family per quota, reproducibly.

    ``exclude`` holds technology ids to leave out, e.g. names the LLM
    judged not to be technologies: they would only teach «rejected».
    """
    excluded = set(exclude)
    by_technology = defaultdict(list)
    for row in rows:
        by_technology[row["technology_id"]].append(row)

    def has_followup(row):
        end = (
            row["horizon_end"]
            if _complete(row, 36)
            else row["horizon_12m_end"]
        )
        return any(
            row["snapshot_date"] < item["date"] <= end
            for later in by_technology[row["technology_id"]]
            if later["snapshot_date"] > row["snapshot_date"]
            for item in _refs(later)
        )

    chosen, used = [], set()
    for group, quota, eligible in quotas:
        families = defaultdict(list)
        for row in rows:
            if row["family_id"] in used or row["technology_id"] in excluded:
                continue
            if eligible(int(row["snapshot_date"][:4]), row):
                families[row["family_id"]].append(row)
        representative = []
        for family, values in families.items():
            followed = [row for row in values if has_followup(row)]
            rich = [row for row in values if int(row["document_count"]) >= 2]
            followed_rich = [row for row in followed if row in rich]
            if followed_rich and int(_rank(group, family)[:2], 16) % 2 == 0:
                candidates = followed_rich
            else:
                candidates = followed or rich or values
            representative.append(
                min(candidates, key=lambda row: row["snapshot_date"])
            )
        order = sorted(
            representative, key=lambda r: _rank(group, r["family_id"])
        )
        followed = [row for row in order if has_followup(row)]
        baseline = [row for row in order if not has_followup(row)]
        count = min(len(followed), round(quota * FOLLOWED_SHARE))
        selected = followed[:count] + baseline[: quota - count]
        if len(selected) < quota:
            selected += followed[count : count + quota - len(selected)]
        for row in selected:
            used.add(row["family_id"])
            chosen.append(
                {
                    **row,
                    "sample_group": group
                    + ("; follow-up" if has_followup(row) else "; baseline"),
                }
            )

    result = []
    for row in sorted(
        chosen, key=lambda r: (r["snapshot_date"], r["technology"])
    ):
        start, year_1, year_3 = (
            row["snapshot_date"],
            row["horizon_12m_end"],
            row["horizon_end"],
        )
        before = {item["document_id"]: item for item in _refs(row)}
        after_1, after_3 = {}, {}
        for later in by_technology[row["technology_id"]]:
            if later["snapshot_date"] <= start:
                continue
            for item in _refs(later):
                if start < item["date"] <= year_1:
                    after_1[item["document_id"]] = item
                elif year_1 < item["date"] <= year_3:
                    after_3[item["document_id"]] = item
        result.append(
            {
                **{
                    key: row.get(key, "")
                    for key in REVIEW_FIELDS
                    if key in row
                },
                "pre_t_sources": show_sources(before),
                "followup_0_12m_sources": show_sources(after_1),
                "followup_13_36m_sources": show_sources(after_3),
                "sample_group": row["sample_group"],
            }
        )
    return result


def add_llm_hints(
    rows: List[Dict[str, Any]], answers: Mapping[str, Mapping[str, Any]]
) -> int:
    """Fill the llm_* columns with the LLM's view of the snapshot's year."""
    filled = 0
    for row in rows:
        answer = answers.get(row["technology_id"])
        if not answer:
            continue
        years = {item["year"]: item for item in answer.get("years") or []}
        year = int(row["snapshot_date"][:4])
        known = [value for key, value in sorted(years.items()) if key <= year]
        at_t = known[-1] if known else (years[min(years)] if years else {})
        row.update(
            {
                "llm_verdict": answer.get("verdict"),
                "llm_score_at_t": at_t.get("score"),
                "llm_hype_at_t": at_t.get("hype"),
                "llm_maturity_at_t": at_t.get("maturity"),
                "llm_rationale": answer.get("rationale"),
            }
        )
        filled += 1
    return filled


def read_list(path: Path) -> List[Dict[str, Any]]:
    """The organisers' 100 signals: from a batch workbook («Список 100»)
    or from the organisers' own file (first sheet, data from row 3)."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    if "Список 100" in book.sheetnames:
        rows = list(book["Список 100"].iter_rows(values_only=True))
        header = [str(cell) for cell in rows[0]]
        return [
            {key: value for key, value in zip(header, row)}
            for row in rows[1:]
            if row and row[0] is not None
        ]
    import re

    result = []
    for record in list(book.active.iter_rows(values_only=True))[2:]:
        if not isinstance(record[1], int):
            continue
        citations = re.findall(
            r"\[[^]]+\]\((https?://[^)]+)\)", record[9] or ""
        )
        result.append(
            {
                "source_no": record[1],
                "technology": record[2],
                "domain": record[3],
                "companies": record[4],
                "why_weak": record[5],
                "stage": record[6],
                "trend": record[7],
                "source_score": record[8],
                "source_urls": "\n".join(citations) or record[9] or "",
            }
        )
    return result


def write_workbook(
    path: Path,
    review: Sequence[Mapping[str, Any]],
    supplied: Sequence[Mapping[str, Any]],
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    guide = book.active
    guide.title = "Правила"
    for row in GUIDE:
        guide.append(list(row))
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 110
    for cells in guide.iter_rows():
        cells[0].font = Font(bold=True)
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    guide["A1"].font = Font(bold=True, size=13)

    fill = PatternFill("solid", fgColor="FFF4CE")
    hint = PatternFill("solid", fgColor="EEF4FC")
    sheet = book.create_sheet("Исторические срезы")
    sheet.append(list(REVIEW_FIELDS))
    for row in review:
        sheet.append([row.get(field) for field in REVIEW_FIELDS])
    last = len(review) + 1
    for index, field in enumerate(REVIEW_FIELDS, 1):
        column = sheet.cell(1, index).column_letter
        sheet.column_dimensions[column].width = (
            60
            if field.endswith("_sources") or field == "llm_rationale"
            else 34
            if field == "technology"
            else 14
        )
        for cell in sheet[column][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if field in REVIEWER_FIELDS:
                cell.fill = fill
            elif field.startswith("llm_"):
                cell.fill = hint
        sheet.cell(1, index).font = Font(bold=True)
    states = DataValidation(type="list", formula1='"' + ",".join(STATES) + '"')
    binary = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
    sheet.add_data_validation(states)
    sheet.add_data_validation(binary)
    states.add(f"D2:D{last}")
    binary.add(f"E2:H{last}")
    sheet.freeze_panes = "D2"

    listed = book.create_sheet("Список 100")
    listed.append(list(LIST_FIELDS))
    for row in supplied:
        listed.append([row.get(field) for field in LIST_FIELDS])
    for index, field in enumerate(LIST_FIELDS, 1):
        column = listed.cell(1, index).column_letter
        listed.column_dimensions[column].width = (
            50
            if field
            in ("technology", "why_weak", "graph_technology", "source_urls")
            else 16
        )
        listed.cell(1, index).font = Font(bold=True)
        for cell in listed[column][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if field in ("match_status", "match_notes"):
                cell.fill = fill
    match = DataValidation(
        type="list", formula1='"same,related,none"', allow_blank=True
    )
    listed.add_data_validation(match)
    match.add(f"M2:M{len(supplied) + 1}")
    listed.freeze_panes = "C2"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def read_annotations(path: Path) -> List[Dict[str, Any]]:
    """Rows of «Исторические срезы» where a person set anything."""
    from openpyxl import load_workbook

    rows = list(
        load_workbook(path, read_only=True, data_only=True)[
            "Исторические срезы"
        ].iter_rows(values_only=True)
    )
    header = [str(cell) for cell in rows[0]]
    result = []
    for values in rows[1:]:
        row = dict(zip(header, values))
        if any(row.get(field) not in (None, "") for field in REVIEWER_FIELDS):
            snapshot = row.get("snapshot_date")
            row["snapshot_date"] = (
                snapshot.date().isoformat()
                if hasattr(snapshot, "date")
                else str(snapshot)[:10]
            )
            result.append(row)
    return result


def main(argv=None) -> Dict[str, Any]:
    parser = argparse.ArgumentParser(
        prog="python -m lctrend.modeling.dataset.annotation_batch",
        description="Build the expert annotation workbook.",
    )
    parser.add_argument("--review", type=Path, required=True)
    parser.add_argument("--llm", type=Path, help="llm_labels.csv.jsonl")
    parser.add_argument("--list", type=Path, help="Organisers' 100 signals")
    parser.add_argument(
        "--matches", type=Path, help="benchmark.csv with graph matches"
    )
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    csv.field_size_limit(1 << 30)
    with args.review.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    answers = {}
    if args.llm:
        from ..labeling.graph_labels import read_answers

        answers = read_answers(args.llm)
    noise = [
        key for key, item in answers.items() if not item.get("is_technology")
    ]
    review = select_batch(rows, exclude=noise)
    hinted = add_llm_hints(review, answers)
    supplied = read_list(args.list) if args.list else []
    if args.matches and supplied:
        with args.matches.open(encoding="utf-8", newline="") as stream:
            matches = {
                str(row["source_no"]): row for row in csv.DictReader(stream)
            }
        for row in supplied:
            found = matches.get(str(row.get("source_no")))
            if found:
                row["graph_technology_id"] = found.get("graph_technology_id")
                row["graph_technology"] = found.get("graph_technology")
                row["match_cosine"] = found.get("match_cosine")
    write_workbook(args.output, review, supplied)
    summary = {
        "rows": len(review),
        "families": len({row["family_id"] for row in review}),
        "with_llm_hint": hinted,
        "excluded_non_technologies": len(noise),
        "list": len(supplied),
        "output": str(args.output),
    }
    print(json.dumps(summary, ensure_ascii=False))
    return summary


if __name__ == "__main__":
    main()
