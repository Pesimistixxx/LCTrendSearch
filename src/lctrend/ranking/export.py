"""The weak-signal table as a file: XLSX (analyst layout), CSV or JSON.

Columns follow the analyst table of September 2026: number, niche, area,
companies, why it is a weak signal, stage, mention trend, score
(stage + trend) and sources as Markdown links.
"""

from __future__ import annotations

import csv
import json
from datetime import date
from pathlib import Path
from typing import Any, List, Mapping

COLUMNS = [
    "№",
    "Технология (слабый сигнал)",
    "Область",
    "Компании",
    "Почему это слабый сигнал",
    "Стадия развития",
    "Тренд упоминаний",
    "Балл (стадия+тренд)",
    "Источники",
]
WIDTHS = [5, 45, 18, 40, 70, 30, 55, 12, 70]
MONTHS = [
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
]


def _link(source: Mapping[str, Any]) -> str:
    title = str(source.get("title") or "").replace("]", ")")
    return f"[{title}]({source['url']})" if source.get("url") else title


def rows(result: Mapping[str, Any]) -> List[List[Any]]:
    return [
        [
            card["number"],
            card["title"],
            card["area"],
            ", ".join(card["companies"]),
            card["why_weak"],
            card["stage_label"],
            card["trend_label"],
            card["score"],
            ", ".join(_link(source) for source in card["sources"]),
        ]
        for card in result["cards"]
    ]


def title(result: Mapping[str, Any]) -> str:
    snapshot = date.fromisoformat(result["snapshot"])
    areas = [area for area in result.get("areas") or [] if area != "Другое"]
    return (
        f"{len(result['cards'])} слабых технологических сигналов "
        f"({MONTHS[snapshot.month - 1]} {snapshot.year})"
        + (": " + ", ".join(areas) if areas else "")
    )


def _xlsx(result: Mapping[str, Any], path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            'XLSX needs openpyxl: pip install -e ".[report]" '
            "(or write .csv / .json)"
        ) from exc
    book = Workbook()
    sheet = book.active
    sheet.title = "Слабые сигналы"
    sheet.append([title(result)])
    sheet.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS)
    )
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(COLUMNS)
    fill = PatternFill("solid", fgColor="DDE7F3")
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows(result):
        sheet.append(row)
    wrap = Alignment(wrap_text=True, vertical="top")
    for line in sheet.iter_rows(min_row=3):
        for cell in line:
            cell.alignment = wrap
    for index, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "C3"
    sheet.auto_filter.ref = (
        f"A2:{get_column_letter(len(COLUMNS))}{sheet.max_row}"
    )
    book.save(path)


def write_report(result: Mapping[str, Any], path: Path) -> Path:
    """Write the table by the file suffix: .xlsx, .csv or .json."""
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        _xlsx(result, path)
    elif suffix == ".csv":
        # utf-8-sig: Excel opens Cyrillic CSV correctly only with a BOM.
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(COLUMNS)
            writer.writerows(rows(result))
    elif suffix == ".json":
        path.write_text(
            json.dumps(result, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    else:
        raise ValueError("report file must end with .xlsx, .csv or .json")
    return path
