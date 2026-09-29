"""Read-only selection of a reproducible, manageable annotation batch."""

import csv
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCE = Path(
    "/Users/komputer/Downloads/100_слабых_технологических_сигналов_сентябрь_2026.xlsx"
)
REVIEW = ROOT / "full-review-2026-09-29.reviewer_1.csv"
QUOTAS = (
    ("1959–2010, 36m", 10, lambda year, row: year <= 2010 and row["calendar_complete_36m"] == "True"),
    ("2011–2016, 36m", 20, lambda year, row: 2011 <= year <= 2016 and row["calendar_complete_36m"] == "True"),
    ("2017–2020, 36m", 100, lambda year, row: 2017 <= year <= 2020 and row["calendar_complete_36m"] == "True"),
    ("2021–2023, 36m", 110, lambda year, row: 2021 <= year <= 2023 and row["calendar_complete_36m"] == "True"),
    ("2024, 12m", 30, lambda year, row: year == 2024 and row["calendar_complete_12m"] == "True"),
    ("2025, 12m", 30, lambda year, row: year == 2025 and row["calendar_complete_12m"] == "True"),
)


def rank(group, row):
    value = f"annotation-batch-2026-09-29:{group}:{row['family_id']}"
    return hashlib.sha256(value.encode()).hexdigest()


def refs(row):
    return json.loads(row["recent_documents"] or "[]")


def show_sources(documents, limit=4):
    ordered = sorted(
        documents.values(),
        key=lambda item: (item["date"], item["document_id"]),
        reverse=True,
    )[:limit]
    return "\n".join(
        f"{item['date']} | {str(item.get('title') or item['document_id'])[:95]} | {item.get('url') or item['document_id']}"
        for item in ordered
    )


def main():
    with REVIEW.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    by_technology = defaultdict(list)
    for row in rows:
        by_technology[row["technology_id"]].append(row)

    def has_followup(row):
        end = (
            row["horizon_end"]
            if row["calendar_complete_36m"] == "True"
            else row["horizon_12m_end"]
        )
        return any(
            row["snapshot_date"] < item["date"] <= end
            for later in by_technology[row["technology_id"]]
            if later["snapshot_date"] > row["snapshot_date"]
            for item in refs(later)
        )

    chosen = []
    used_families = set()
    for group, quota, eligible in QUOTAS:
        families = defaultdict(list)
        for row in rows:
            family = row["family_id"]
            if family in used_families:
                continue
            if eligible(int(row["snapshot_date"][:4]), row):
                families[family].append(row)
        representative = []
        for values in families.values():
            followed = [row for row in values if has_followup(row)]
            rich_values = [row for row in values if int(row["document_count"]) >= 2]
            followed_rich = [row for row in followed if int(row["document_count"]) >= 2]
            if followed_rich and int(rank(group, values[0])[:2], 16) % 2 == 0:
                candidates = followed_rich
            else:
                candidates = followed or rich_values or values
            representative.append(
                min(candidates, key=lambda row: row["snapshot_date"])
            )
        followed = sorted(
            (row for row in representative if has_followup(row)),
            key=lambda row: rank(group, row),
        )
        baseline = sorted(
            (row for row in representative if not has_followup(row)),
            key=lambda row: rank(group, row),
        )
        followed_count = min(len(followed), round(quota * 0.6))
        selected = followed[:followed_count] + baseline[: quota - followed_count]
        if len(selected) < quota:
            selected += followed[followed_count : followed_count + quota - len(selected)]
        assert len(selected) == quota, (group, len(selected), quota)
        for row in selected:
            used_families.add(row["family_id"])
            selected_row = dict(row)
            selected_row["sample_group"] = group + (
                "; follow-up" if has_followup(row) else "; baseline"
            )
            chosen.append(selected_row)

    review = []
    for row in sorted(chosen, key=lambda item: (item["snapshot_date"], item["technology"])):
        start = row["snapshot_date"]
        year_1 = row["horizon_12m_end"]
        year_3 = row["horizon_end"]
        before = {item["document_id"]: item for item in refs(row)}
        after_1, after_3 = {}, {}
        for followup in by_technology[row["technology_id"]]:
            if followup["snapshot_date"] <= start:
                continue
            for item in refs(followup):
                when = item["date"]
                if start < when <= year_1:
                    after_1[item["document_id"]] = item
                elif year_1 < when <= year_3:
                    after_3[item["document_id"]] = item
        review.append(
            {
                **{key: row[key] for key in (
                    "technology_id", "technology", "family_id", "snapshot_date",
                    "horizon_12m_end", "horizon_end", "calendar_complete_12m",
                    "calendar_complete_36m", "document_count", "first_seen_date",
                    "source_families", "organizations", "reviewer_state",
                    "reviewer_signal_12m", "reviewer_trend_12m",
                    "reviewer_signal_36m", "reviewer_trend_36m", "evidence_notes",
                )},
                "pre_t_sources": show_sources(before),
                "followup_0_12m_sources": show_sources(after_1),
                "followup_13_36m_sources": show_sources(after_3),
                "sample_group": row["sample_group"],
            }
        )

    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    supplied = []
    for record in list(workbook.active.iter_rows(values_only=True))[2:]:
        if not isinstance(record[1], int):
            continue
        citations = re.findall(r"\[[^]]+\]\((https?://[^)]+)\)", record[9] or "")
        supplied.append(
            {
                "source_no": record[1],
                "technology": record[2],
                "domain": record[3],
                "companies": record[4],
                "why_weak": record[5],
                "stage": record[6],
                "trend": record[7],
                "source_score": record[8],
                "source_urls": "\n".join(citations) or record[9] or "",
                "graph_technology_id": "",
                "match_status": "",
                "match_notes": "",
            }
        )
    assert len(review) == 300 and len(supplied) == 100
    assert len({(r["technology_id"], r["snapshot_date"]) for r in review}) == 300
    assert len({r["family_id"] for r in review}) == 300
    print(json.dumps({"review": review, "supplied": supplied}, ensure_ascii=False))


if __name__ == "__main__":
    main()
