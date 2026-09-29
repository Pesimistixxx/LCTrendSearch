import json

import pytest

from lctrend.modeling.dataset.annotation_batch import (
    add_llm_hints,
    read_annotations,
    select_batch,
    write_workbook,
)
from lctrend.modeling.labeling.benchmark import attach_scores, nearest, summary


def _row(technology, family, when, documents=2, later_refs=()):
    return {
        "technology_id": technology,
        "technology": technology.upper(),
        "family_id": family,
        "snapshot_date": when,
        "horizon_12m_end": when.replace(when[:4], str(int(when[:4]) + 1)),
        "horizon_end": when.replace(when[:4], str(int(when[:4]) + 3)),
        "calendar_complete_12m": "True",
        "calendar_complete_36m": "True",
        "document_count": str(documents),
        "first_seen_date": when,
        "source_families": "scholarly",
        "organizations": "",
        "recent_documents": json.dumps(
            [
                {
                    "document_id": f"{technology}-{d}",
                    "date": d,
                    "title": f"on {technology}",
                }
                for d in later_refs
            ]
        ),
    }


def test_batch_takes_one_row_per_family_within_quotas_and_skips_noise():
    rows = [
        _row("a", "f1", "2018-01-01"),
        _row("a", "f1", "2019-01-01", later_refs=("2019-06-01",)),
        _row("b", "f1", "2018-07-01"),
        _row("c", "f2", "2018-01-01"),
        _row("noise", "f3", "2018-01-01"),
    ]
    quotas = (("2017–2020", 2, lambda year, row: 2017 <= year <= 2020),)
    batch = select_batch(rows, quotas, exclude=["noise"])
    assert len(batch) == 2
    assert {row["family_id"] for row in batch} == {"f1", "f2"}
    assert all(row["technology_id"] != "noise" for row in batch)
    assert all(row["sample_group"].startswith("2017–2020") for row in batch)


def test_llm_hint_is_the_view_of_the_snapshot_year():
    batch = [{"technology_id": "a", "snapshot_date": "2020-04-01"}]
    answers = {
        "a": {
            "verdict": "success",
            "rationale": "grew",
            "years": [
                {"year": 2019, "score": 0.1, "hype": 0.2, "maturity": 0.3},
                {"year": 2020, "score": 0.2, "hype": 0.3, "maturity": 0.4},
                {"year": 2024, "score": 1.0, "hype": 0.1, "maturity": 0.9},
            ],
        }
    }
    assert add_llm_hints(batch, answers) == 1
    assert batch[0]["llm_score_at_t"] == 0.2
    assert batch[0]["llm_verdict"] == "success"


def test_workbook_round_trip_returns_only_filled_rows(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    review = [
        {
            "technology_id": "a",
            "technology": "A",
            "snapshot_date": "2020-01-01",
        },
        {
            "technology_id": "b",
            "technology": "B",
            "snapshot_date": "2021-01-01",
        },
    ]
    path = write_workbook(
        tmp_path / "batch.xlsx",
        review,
        [{"source_no": 1, "technology": "Сигнал", "match_cosine": 0.9}],
    )
    book = load_workbook(path)
    assert book.sheetnames == ["Правила", "Исторические срезы", "Список 100"]
    sheet = book["Исторические срезы"]
    header = [cell.value for cell in sheet[1]]
    sheet.cell(2, header.index("reviewer_state") + 1, "weak")
    sheet.cell(2, header.index("reviewer_signal_36m") + 1, 1)
    book.save(path)
    labels = read_annotations(path)
    assert [
        (row["technology_id"], row["reviewer_state"]) for row in labels
    ] == [("a", "weak")]
    assert labels[0]["snapshot_date"] == "2020-01-01"


def test_benchmark_matches_by_cosine_and_measures_recall():
    concepts = [
        {"concept_id": "x", "label": "X", "vector": [1.0, 0.0]},
        {"concept_id": "y", "label": "Y", "vector": [0.0, 1.0]},
    ]
    found = nearest([[0.9, 0.1], [0.1, 0.9]], concepts, top=2)
    assert [items[0]["concept_id"] for items in found] == ["x", "y"]
    matches = [
        {"source_no": 1, "graph_technology_id": "x", "match_cosine": 0.95},
        {"source_no": 2, "graph_technology_id": "y", "match_cosine": 0.5},
    ]
    rows = attach_scores(
        matches,
        [{"technology_id": "x", "probability": "0.8", "signal": "1"}],
        {
            "x": {
                "verdict": "niche",
                "is_technology": True,
                "years": [{"year": 2026, "score": 0.2}],
            }
        },
    )
    result = summary(rows, min_cosine=0.8)
    assert result["matched"] == 1
    assert result["model_recall"] == 1.0
    assert result["llm_recall"] == 1.0
